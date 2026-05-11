import reflex as rx
from typing import TypedDict, Optional
from datetime import datetime
from app.states.auth_state import AuthState
from app.shared_store import SharedStore


class Document(TypedDict):
    id: str
    name: str
    category: str
    date: str
    size: str
    path: str


class Shipment(TypedDict):
    id: str
    origin: str
    destination: str
    status: str
    progress: int
    eta: str
    items: int
    priority: str
    destinatario: str
    direccion: str
    valoracion: str
    guia: str
    mes: str
    docs_obs: str
    observaciones: str


class DocuState(rx.State):
    active_tab: str = "documentos"
    is_sidebar_open: bool = False
    documents: list[Document] = []
    shipments: list[Shipment] = []
    selected_docs: list[str] = []
    search_tags: list[str] = []
    input_value: str = ""
    open_menu_id: str = ""
    custom_categories: list[str] = []
    new_category_input: str = ""
    drive_link_input: str = ""
    is_importing: bool = False
    import_message: str = ""
    shipment_search_query: str = ""
    total_documents: int = 0
    active_shipments: int = 0
    delivered_shipments: int = 0
    total_shipments_count: int = 0
    delivered_percentage: int = 0
    in_transit_percentage: int = 0
    pending_percentage: int = 0
    total_units: int = 0
    total_value: float = 0.0
    status_chart_data: list[dict[str, str | int]] = []
    monthly_chart_data: list[dict[str, str | int]] = []
    top_origins: list[dict[str, str | int | float]] = []
    top_destinations: list[dict[str, str | int | float]] = []
    recent_shipments_list: list[dict] = []
    unique_categories: list[dict[str, str | int]] = []
    recent_documents: list[dict] = []
    filtered_shipments: list[dict] = []
    filtered_documents: list[dict] = []

    def _recalculate_filters(self):
        """Helper to recalculate filtered results based on search state."""
        if not self.shipment_search_query.strip():
            self.filtered_shipments = self.shipments
        else:
            query = self.shipment_search_query.lower().strip()
            self.filtered_shipments = [
                s
                for s in self.shipments
                if query in s["id"].lower() or query in s["guia"].lower()
            ]
        docs = self.documents
        search_terms = [tag.lower() for tag in self.search_tags]
        if self.input_value.strip():
            search_terms.append(self.input_value.lower())
        if not search_terms:
            self.filtered_documents = docs
        else:
            filtered = []
            for doc in docs:
                match = False
                for term in search_terms:
                    if (
                        term in doc["name"].lower()
                        or term in doc["category"].lower()
                    ):
                        match = True
                        break
                if match:
                    filtered.append(doc)
            self.filtered_documents = filtered

    def _recalculate_stats(self):
        """Helper to recalculate all stats and derived data."""
        self.total_documents = len(self.documents)
        self.total_shipments_count = len(self.shipments)
        self.active_shipments = len(
            [s for s in self.shipments if s["status"] != "Entregado"]
        )
        self.delivered_shipments = len(
            [s for s in self.shipments if s["status"] == "Entregado"]
        )
        if self.shipments:
            total = len(self.shipments)
            self.delivered_percentage = int(
                round(self.delivered_shipments / total * 100)
            )
            self.in_transit_percentage = int(
                round(
                    sum(
                        (
                            1
                            for s in self.shipments
                            if s["status"] == "En camino"
                        )
                    )
                    / total
                    * 100
                )
            )
            self.pending_percentage = int(
                round(
                    sum(
                        (
                            1
                            for s in self.shipments
                            if s["status"] == "Pendiente"
                        )
                    )
                    / total
                    * 100
                )
            )
        else:
            self.delivered_percentage = 0
            self.in_transit_percentage = 0
            self.pending_percentage = 0
        self.total_units = sum((s.get("items", 0) for s in self.shipments))
        val_sum = 0.0
        for s in self.shipments:
            try:
                val_sum += float(s.get("valoracion", 0) or 0)
            except (ValueError, TypeError):
                pass
        self.total_value = val_sum
        counts = {"Entregado": 0, "En camino": 0, "Pendiente": 0}
        for s in self.shipments:
            st = s.get("status")
            if st in counts:
                counts[st] += 1
        self.status_chart_data = [
            {
                "name": "Entregado",
                "value": counts["Entregado"],
                "fill": "#22c55e",
            },
            {
                "name": "En camino",
                "value": counts["En camino"],
                "fill": "#3b82f6",
            },
            {
                "name": "Pendiente",
                "value": counts["Pendiente"],
                "fill": "#f59e0b",
            },
        ]
        month_map = {
            "1": "Ene",
            "2": "Feb",
            "3": "Mar",
            "4": "Abr",
            "5": "May",
            "6": "Jun",
            "7": "Jul",
            "8": "Ago",
            "9": "Sep",
            "10": "Oct",
            "11": "Nov",
            "12": "Dic",
        }
        monthly_data = {}
        for s in self.shipments:
            m_name = month_map.get(s.get("mes", ""), s.get("mes", "N/A"))
            if m_name not in monthly_data:
                monthly_data[m_name] = {"total": 0, "entregados": 0}
            monthly_data[m_name]["total"] += 1
            if s.get("status") == "Entregado":
                monthly_data[m_name]["entregados"] += 1
        self.monthly_chart_data = [
            {"mes": m, "total": d["total"], "entregados": d["entregados"]}
            for m, d in monthly_data.items()
        ]
        o_counts, d_counts = ({}, {})
        for s in self.shipments:
            o, d = (s.get("origin", ""), s.get("destination", ""))
            if o:
                o_counts[o] = o_counts.get(o, 0) + 1
            if d:
                d_counts[d] = d_counts.get(d, 0) + 1
        sorted_o = sorted(o_counts.items(), key=lambda x: x[1], reverse=True)[
            :5
        ]
        max_o = max([c for _, c in sorted_o]) if sorted_o else 1
        self.top_origins = [
            {"city": city, "count": count, "width": f"{count / max_o * 100}%"}
            for city, count in sorted_o
        ]
        sorted_d = sorted(d_counts.items(), key=lambda x: x[1], reverse=True)[
            :5
        ]
        max_d = max([c for _, c in sorted_d]) if sorted_d else 1
        self.top_destinations = [
            {"city": city, "count": count, "width": f"{count / max_d * 100}%"}
            for city, count in sorted_d
        ]
        self.recent_shipments_list = [dict(s) for s in self.shipments[:10]]
        self.recent_documents = [dict(d) for d in self.documents[:5]]
        cats = {doc["category"]: 0 for doc in self.documents}
        for doc in self.documents:
            cats[doc["category"]] += 1
        for custom in self.custom_categories:
            if custom not in cats:
                cats[custom] = 0
        self.unique_categories = [
            {"name": n, "count": c} for n, c in cats.items()
        ]
        self._recalculate_filters()

    @rx.event
    def load_shared_data(self):
        data = SharedStore.load_all_data()
        self.documents = data["documents"]
        self.shipments = data["shipments"]
        self.custom_categories = data["custom_categories"]
        self._recalculate_stats()

    @rx.event
    def set_tab(self, tab: str):
        self.active_tab = tab
        self.is_sidebar_open = False

    @rx.event
    def toggle_sidebar(self):
        self.is_sidebar_open = not self.is_sidebar_open

    @rx.event
    def set_input_value(self, val: str):
        self.input_value = val
        self._recalculate_filters()

    @rx.event
    def set_shipment_search_query(self, val: str):
        self.shipment_search_query = val
        self._recalculate_filters()

    @rx.event
    def toggle_doc_selection(self, doc_id: str):
        if doc_id in self.selected_docs:
            self.selected_docs.remove(doc_id)
        else:
            self.selected_docs.append(doc_id)

    @rx.event
    def toggle_select_all(self):
        filtered_ids = [doc["id"] for doc in self.filtered_documents]
        if (
            len(self.selected_docs) == len(filtered_ids)
            and len(filtered_ids) > 0
        ):
            self.selected_docs = []
        else:
            self.selected_docs = filtered_ids

    @rx.event
    def add_search_tag(self, key_code: str):
        if key_code == "Enter" and self.input_value.strip():
            tag = self.input_value.strip()
            if tag not in self.search_tags:
                self.search_tags.append(tag)
            self.input_value = ""
            self._recalculate_filters()
        elif (
            key_code == "Backspace"
            and (not self.input_value)
            and self.search_tags
        ):
            self.search_tags.pop()
            self._recalculate_filters()

    @rx.event
    def remove_tag(self, tag: str):
        self.search_tags.remove(tag)
        self._recalculate_filters()

    @rx.event
    def delete_document(self, doc_id: str):
        self.documents = [d for d in self.documents if d["id"] != doc_id]
        SharedStore.remove_document(doc_id)
        if doc_id in self.selected_docs:
            self.selected_docs.remove(doc_id)
        self._recalculate_stats()
        return rx.toast("Documento eliminado")

    @rx.event
    def add_custom_category(self):
        cat = self.new_category_input.strip()
        if cat and cat not in self.custom_categories:
            self.custom_categories.append(cat)
            SharedStore.add_custom_category(cat)
        self.new_category_input = ""
        self._recalculate_stats()

    @rx.event
    def clear_search(self):
        self.search_tags = []
        self.input_value = ""
        self._recalculate_filters()

    @rx.event
    def clear_shipment_search(self):
        self.shipment_search_query = ""
        self._recalculate_filters()

    @rx.event
    async def handle_upload(self, files: list[rx.UploadFile]):
        auth = await self.get_state(AuthState)
        if auth.current_role != "admin":
            yield rx.toast("No tienes permisos para subir archivos.")
            return
        import random
        import string

        upload_dir = rx.get_upload_dir()
        upload_dir.mkdir(parents=True, exist_ok=True)
        for file in files:
            upload_data = await file.read()
            random_string = "".join(
                random.choices(string.ascii_letters + string.digits, k=10)
            )
            unique_filename = f"{random_string}_{file.name}"
            file_path = upload_dir / unique_filename
            with file_path.open("wb") as f:
                f.write(upload_data)
            size_mb = f"{len(upload_data) / (1024 * 1024):.2f} MB"
            new_doc: Document = {
                "id": "".join(
                    random.choices(string.ascii_letters + string.digits, k=8)
                ),
                "name": file.name,
                "category": "Subido",
                "date": datetime.now().strftime("%Y-%m-%d"),
                "size": size_mb,
                "path": unique_filename,
            }
            self.documents.insert(0, new_doc)
            SharedStore.add_document(new_doc)
        self._recalculate_stats()
        yield rx.toast("Archivos subidos con éxito")
        return

    @rx.event
    def download_doc(self, doc_id: str):
        doc = next((d for d in self.documents if d["id"] == doc_id), None)
        if doc:
            if doc.get("path"):
                return rx.download(
                    url=f"/_upload/{doc['path']}", filename=doc["name"]
                )
            else:
                upload_dir = rx.get_upload_dir()
                upload_dir.mkdir(parents=True, exist_ok=True)
                content = f"Document: {doc['name']}\nCategory: {doc['category']}\nDate: {doc['date']}\nSize: {doc['size']}\n\nThis is a generated placeholder for the demo document."
                filename = f"demo_{doc['id']}_{doc['name']}.txt"
                file_path = upload_dir / filename
                with file_path.open("wb") as f:
                    f.write(content.encode("utf-8"))
                for d in self.documents:
                    if d["id"] == doc_id:
                        d["path"] = filename
                        break
                SharedStore.update_document(doc_id, {"path": filename})
                return rx.download(
                    url=f"/_upload/{filename}", filename=doc["name"]
                )
        return rx.toast("Documento no encontrado")

    @rx.event
    def download_selected(self):
        if not self.selected_docs:
            return
        events = []
        for doc_id in self.selected_docs:
            events.append(DocuState.download_doc(doc_id))
        return events

    @rx.event
    async def delete_selected_docs(self):
        auth = await self.get_state(AuthState)
        if auth.current_role != "admin":
            yield rx.toast("No tienes permisos para eliminar documentos.")
            return
        count = len(self.selected_docs)
        if count == 0:
            return
        for doc_id in self.selected_docs:
            SharedStore.remove_document(doc_id)
        self.documents = [
            d for d in self.documents if d["id"] not in self.selected_docs
        ]
        self.selected_docs = []
        self._recalculate_stats()
        yield rx.toast(f"{count} documentos eliminados correctamente.")

    @rx.event
    def toggle_shipment_menu(self, ship_id: str):
        self.open_menu_id = "" if self.open_menu_id == ship_id else ship_id

    @rx.event
    def delete_shipment(self, ship_id: str):
        self.shipments = [s for s in self.shipments if s["id"] != ship_id]
        SharedStore.remove_shipment(ship_id)
        self.open_menu_id = ""
        self._recalculate_stats()
        return rx.toast(f"Envío {ship_id} eliminado")

    @rx.event
    def download_project_zip(self):
        return rx.download(
            url="/_upload/trazo_project.zip", filename="trazo_project.zip"
        )

    def _process_rows(self, rows):
        new_shipments = []
        for row in rows:
            if not row or not row[0]:
                continue
            row_data = list(row) + [""] * (12 - len(row))
            guia_raw = str(row_data[0]).strip()
            if guia_raw.endswith(".0"):
                guia_raw = guia_raw[:-2]
            estado_raw = str(row_data[3]).strip()
            if estado_raw == "Entregada":
                status = "Entregado"
                progress = 100
            elif estado_raw in ["En Recibo", "En Traslado Nacional"]:
                status = "En camino"
                progress = 50
            else:
                status = "Pendiente"
                progress = 0
            try:
                items_val = int(float(str(row_data[9]).strip() or 0))
            except ValueError:
                items_val = 1
            ship: Shipment = {
                "id": f"G-{guia_raw}",
                "origin": str(row_data[4]).strip(),
                "destination": f"{str(row_data[5]).strip()}, {str(row_data[7]).strip()}",
                "status": status,
                "progress": progress,
                "eta": str(row_data[2]).strip(),
                "items": items_val,
                "priority": "Normal",
                "destinatario": str(row_data[6]).strip(),
                "direccion": str(row_data[7]).strip(),
                "valoracion": str(row_data[8]).strip(),
                "guia": guia_raw,
                "mes": str(row_data[1]).strip(),
                "docs_obs": str(row_data[10]).strip(),
                "observaciones": str(row_data[11]).strip(),
            }
            new_shipments.append(ship)
        self.shipments = new_shipments

    @rx.event
    async def handle_excel_upload(self, files: list[rx.UploadFile]):
        auth = await self.get_state(AuthState)
        if auth.current_role != "admin":
            yield rx.toast("No tienes permisos para importar datos.")
            return
        if not files:
            return
        self.is_importing = True
        yield
        try:
            import openpyxl
            import io

            file = files[0]
            upload_data = await file.read()
            wb = openpyxl.load_workbook(
                filename=io.BytesIO(upload_data), data_only=True
            )
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) > 1:
                self._process_rows(rows[1:])
                SharedStore.set_shipments(self.shipments)
                self._recalculate_stats()
                self.import_message = f"{len(rows) - 1} envíos importados correctamente desde Excel."
                yield rx.toast(self.import_message)
                self.active_tab = "seguimiento_v1"
            else:
                self.import_message = "El archivo Excel está vacío."
                yield rx.toast(self.import_message)
        except Exception as e:
            import logging

            logging.exception(f"Error procesando Excel: {e}")
            self.import_message = "Error al procesar el archivo Excel."
            yield rx.toast(self.import_message)
        finally:
            self.is_importing = False

    @rx.event
    def import_from_drive(self):
        import re
        import csv
        import urllib.request
        from io import StringIO
        import logging

        self.is_importing = True
        yield
        try:
            match = re.search("/d/([a-zA-Z0-9-_]+)", self.drive_link_input)
            if not match:
                self.import_message = "Enlace de Google Drive inválido."
                yield rx.toast(self.import_message)
                return
            sheet_id = match.group(1)
            csv_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv"
            req = urllib.request.Request(csv_url)
            with urllib.request.urlopen(req) as response:
                csv_text = response.read().decode("utf-8")
            reader = csv.reader(StringIO(csv_text))
            rows = list(reader)
            if len(rows) > 1:
                self._process_rows(rows[1:])
                SharedStore.set_shipments(self.shipments)
                self._recalculate_stats()
                self.import_message = f"{len(rows) - 1} envíos importados correctamente desde Google Drive."
                yield rx.toast(self.import_message)
                self.active_tab = "seguimiento_v1"
            else:
                self.import_message = "El documento está vacío."
                yield rx.toast(self.import_message)
        except Exception as e:
            logging.exception(f"Error importando de Drive: {e}")
            self.import_message = "Error al importar desde Google Drive."
            yield rx.toast(self.import_message)
        finally:
            self.is_importing = False